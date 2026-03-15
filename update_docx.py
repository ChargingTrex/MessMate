"""
MessMate Word Document Updater (update_docx.py)
-----------------------------------------------
A utility script used during development to programmatically find and replace 
strings inside the MessMate_Checklist.docx file using the python-docx library.
"""

from docx import Document
import os

def update_docx():
    file_path = "MessMate_Checklist.docx"
    if not os.path.exists(file_path):
        print(f"Error: {file_path} not found.")
        return

    doc = Document(file_path)
    changes_made = 0

    # Function to replace text in a run, preserving formatting
    def replace_in_element(element):
        nonlocal changes_made
        if 'Chart.js' in element.text:
            element.text = element.text.replace('Chart.js', 'Matplotlib')
            changes_made += 1
        if 'static/dashboard.js' in element.text:
            element.text = element.text.replace('static/dashboard.js', 'Python Matplotlib Rendering')
            changes_made += 1

    # Check paragraphs
    for p in doc.paragraphs:
        for run in p.runs:
            replace_in_element(run)

    # Check tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        replace_in_element(run)

    if changes_made > 0:
        doc.save(file_path)
        print(f"✅ Successfully updated {file_path}. Made {changes_made} replacements.")
    else:
        print("No occurrences of 'Chart.js' or 'static/dashboard.js' found. File not modified.")

if __name__ == "__main__":
    update_docx()
